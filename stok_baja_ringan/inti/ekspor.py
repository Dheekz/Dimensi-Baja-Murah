"""Membuat file Excel laporan dengan susunan sheet seperti file lama:
Real Time, Produksi, Pengeluaran, Barang masuk.
"""
from io import BytesIO
from operator import sub as kurang

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from inti import laporan as lp
from inti import stok as sk
from inti.tema import BAJA_GELAP, OREN, OREN_MUDA, SENG, tanpa_pagar

_GARIS = Side(style="thin", color="BFC5CC")
_BINGKAI = Border(left=_GARIS, right=_GARIS, top=_GARIS, bottom=_GARIS)
_KEPALA = PatternFill(fill_type="solid", start_color=tanpa_pagar(OREN), end_color=tanpa_pagar(OREN))
_KEPALA_GELAP = PatternFill(fill_type="solid", start_color=tanpa_pagar(BAJA_GELAP), end_color=tanpa_pagar(BAJA_GELAP))
_SUBTOTAL = PatternFill(fill_type="solid", start_color=tanpa_pagar(OREN_MUDA), end_color=tanpa_pagar(OREN_MUDA))
_ABU = PatternFill(fill_type="solid", start_color=tanpa_pagar(SENG), end_color=tanpa_pagar(SENG))
_TENGAH = Alignment(horizontal="center", vertical="center", wrap_text=True)
_FORMAT_ANGKA = "#,##0;[Red]\\(#,##0\\);0"

KOLOM_REAL_TIME = [
    ("Supplier", "supplier", 10),
    ("Nama Barang", "nama_barang", 26),
    ("Warna / Ukuran", "varian", 14),
    ("Stok Awal", "stok_awal", 11),
    ("Tgl Stok Awal", "tanggal_stok_awal", 13),
    ("Produksi", "produksi", 11),
    ("Barang Masuk", "barang_masuk", 11),
    ("Dipakai Produksi", "dipakai_produksi", 11),
    ("Pengeluaran", "pengeluaran", 12),
    ("Koreksi Tambah", "koreksi_tambah", 11),
    ("Koreksi Kurang", "koreksi_kurang", 11),
    ("STOK RT", "stok_rt", 11),
    ("Selisih", "selisih", 10),
    ("Status", "status", 15),
]


def _kepala(sel, gelap=False):
    sel.fill = _KEPALA_GELAP if gelap else _KEPALA
    sel.font = Font(bold=True, color="FFFFFF")
    sel.alignment = _TENGAH
    sel.border = _BINGKAI


def tulis_real_time(ws, data_stok):
    for i, (judul, _, lebar) in enumerate(KOLOM_REAL_TIME, start=1):
        _kepala(ws.cell(row=1, column=i, value=judul), gelap=(judul == "STOK RT"))
        ws.column_dimensions[get_column_letter(i)].width = lebar
    ws.row_dimensions[1].height = 32
    baris_ke = 2
    for tipe in (sk.TIPE_JADI, sk.TIPE_BAHAN):
        bagian = data_stok[data_stok["tipe"] == tipe]
        if bagian.empty:
            continue
        sel = ws.cell(row=baris_ke, column=1, value=tipe)
        sel.font = Font(bold=True)
        for c in range(1, len(KOLOM_REAL_TIME) + 1):
            ws.cell(row=baris_ke, column=c).fill = _ABU
        baris_ke += 1
        for (_, nama), kelompok in bagian.groupby(["supplier", "nama_barang"], sort=False):
            awal_kelompok = baris_ke
            for _, item in kelompok.iterrows():
                for c, (_, kunci, _) in enumerate(KOLOM_REAL_TIME, start=1):
                    nilai = item[kunci]
                    if kunci == "tanggal_stok_awal":
                        nilai = nilai.to_pydatetime() if pd.notna(nilai) else None
                    sel = ws.cell(row=baris_ke, column=c, value=nilai)
                    sel.border = _BINGKAI
                    if kunci == "tanggal_stok_awal":
                        sel.number_format = "dd/mm/yyyy"
                    elif isinstance(nilai, (int, float)):
                        sel.number_format = _FORMAT_ANGKA
                    if kunci == "stok_rt":
                        sel.font = Font(bold=True)
                        sel.fill = _SUBTOTAL
                    if kunci == "status" and nilai in sk.STATUS_PERHATIAN:
                        sel.font = Font(bold=True, color="B83227")
                baris_ke += 1
            if len(kelompok) > 1:
                for c in (1, 2):
                    ws.merge_cells(start_row=awal_kelompok, start_column=c, end_row=kurang(baris_ke, 1), end_column=c)
                    ws.cell(row=awal_kelompok, column=c).alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "D2"


def tulis_mingguan(ws, tabel, kolom_meta):
    info = lp.KOLOM_INFO
    lebar_info = [10, 26, 14]
    for i, (judul, lebar) in enumerate(zip(info, lebar_info), start=1):
        ws.merge_cells(start_row=1, start_column=i, end_row=2, end_column=i)
        _kepala(ws.cell(row=1, column=i, value=judul))
        _kepala(ws.cell(row=2, column=i))
        ws.column_dimensions[get_column_letter(i)].width = lebar
    for j, k in enumerate(kolom_meta, start=len(info) + 1):
        huruf = get_column_letter(j)
        ws.column_dimensions[huruf].width = 11 if k["total"] else 9
        if k["total"]:
            ws.merge_cells(start_row=1, start_column=j, end_row=2, end_column=j)
            _kepala(ws.cell(row=1, column=j, value=k["hari"]), gelap=True)
            _kepala(ws.cell(row=2, column=j), gelap=True)
        else:
            _kepala(ws.cell(row=1, column=j, value=k["hari"]))
            sel = ws.cell(row=2, column=j, value=k["tanggal"].to_pydatetime())
            _kepala(sel)
            sel.number_format = "dd/mm/yyyy"
    ws.row_dimensions[1].height = 30

    semua_kunci = info + [k["kunci"] for k in kolom_meta]
    kolom_total = [False] * len(info) + [k["total"] for k in kolom_meta]
    for r, (_, baris) in enumerate(tabel.iterrows(), start=3):
        jenis_baris = baris[lp.PENANDA]
        for c, (kunci, adalah_total) in enumerate(zip(semua_kunci, kolom_total), start=1):
            nilai = baris[kunci]
            if isinstance(nilai, float):
                nilai = int(nilai) if nilai.is_integer() else nilai
                if nilai == 0 and jenis_baris == "isi":
                    nilai = None
            sel = ws.cell(row=r, column=c, value=nilai)
            sel.border = _BINGKAI
            if c > len(info):
                sel.number_format = _FORMAT_ANGKA
                sel.alignment = Alignment(horizontal="center")
            if jenis_baris == "subtotal":
                sel.fill = _SUBTOTAL
                sel.font = Font(bold=True)
            elif jenis_baris == "total":
                sel.fill = _KEPALA_GELAP
                sel.font = Font(bold=True, color="FFFFFF")
            elif adalah_total:
                sel.fill = _ABU
                sel.font = Font(bold=True)
    ws.freeze_panes = ws.cell(row=3, column=len(info) + 1)


def buat_excel(master, transaksi, kategori_list, mulai, selesai):
    """Satu file Excel. Bila kategori lebih dari satu, nama sheet diberi akhiran kategori."""
    buku = Workbook()
    buku.remove(buku.active)
    data_stok = sk.hitung_stok(master, transaksi, sampai=selesai)
    banyak = len(kategori_list) > 1
    for kategori in kategori_list:
        akhiran = f" {kategori}"[:12] if banyak else ""
        tulis_real_time(buku.create_sheet(f"Real Time{akhiran}"[:31]), data_stok[data_stok["kategori"] == kategori])
        for judul, jenis in (("Produksi", sk.PRODUKSI), ("Pengeluaran", sk.PENGELUARAN), ("Barang masuk", sk.BARANG_MASUK)):
            tabel, meta = lp.tabel_mingguan(master, transaksi, kategori, jenis, mulai, selesai)
            tulis_mingguan(buku.create_sheet(f"{judul}{akhiran}"[:31]), tabel, meta)
    wadah = BytesIO()
    buku.save(wadah)
    return wadah.getvalue()


def buat_excel_tabel(df, nama_sheet="Data"):
    """Ekspor tabel sederhana (misal riwayat transaksi) ke Excel."""
    buku = Workbook()
    ws = buku.active
    ws.title = nama_sheet[:31]
    for c, judul in enumerate(df.columns, start=1):
        _kepala(ws.cell(row=1, column=c, value=str(judul)))
        ws.column_dimensions[get_column_letter(c)].width = max(10, min(40, len(str(judul)) + 4))
    for r, baris in enumerate(df.itertuples(index=False), start=2):
        for c, nilai in enumerate(baris, start=1):
            if isinstance(nilai, pd.Timestamp):
                nilai = nilai.to_pydatetime() if pd.notna(nilai) else None
            sel = ws.cell(row=r, column=c, value=nilai)
            sel.border = _BINGKAI
            if hasattr(nilai, "strftime"):
                sel.number_format = "dd/mm/yyyy"
    ws.freeze_panes = "A2"
    wadah = BytesIO()
    buku.save(wadah)
    return wadah.getvalue()
